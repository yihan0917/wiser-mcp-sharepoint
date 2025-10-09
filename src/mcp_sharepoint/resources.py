"""
SharePoint resources using Microsoft Graph API
"""
import base64, os, fitz, io, logging, time, pandas as pd
import requests
from typing import Dict, Any, List, Optional
from docx import Document
from pptx import Presentation
from .common import logger, ACCESS_TOKEN, SITE_ID, DRIVE_ID, make_graph_request, get_graph_headers

logger = logging.getLogger(__name__)

def extract_text_from_pdf(pdf_content):
    """Extract text from PDF using PyMuPDF"""
    try:
        pdf_document = fitz.open(stream=pdf_content, filetype="pdf")
        text_content = "".join(pdf_document[i].get_text() + "\n" for i in range(len(pdf_document)))
        page_count = len(pdf_document)
        pdf_document.close()
        return text_content.strip(), page_count
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {e}")
        raise

# def extract_text_from_excel(content_bytes):
#     """Extract text from Excel files"""
#     try:
#         sheets = pd.read_excel(io.BytesIO(content_bytes), sheet_name=None)
#         text_parts = []
#         for sheet_name, df in sheets.items():
#             text_parts.append(f"=== {sheet_name} ===")
#             text_parts.extend(df.head(50).fillna('').astype(str).apply(' | '.join, axis=1).tolist())
#         return "\n".join(text_parts), len(sheets)
#     except Exception as e:
#         logger.error(f"Error extracting text from Excel: {e}")
#         raise

def extract_text_from_excel(content_bytes, max_rows_per_sheet=None):
    """Extract text from Excel files"""
    try:
        sheets = pd.read_excel(io.BytesIO(content_bytes), sheet_name=None)
        text_parts = []
        for sheet_name, df in sheets.items():
            text_parts.append(f"=== {sheet_name} ===")
            
            # Add column headers
            if not df.empty:
                headers = " | ".join(str(col) for col in df.columns)
                text_parts.append(f"HEADERS: {headers}")
                text_parts.append("-" * len(headers))  # Separator line
                
                # Add data rows
                # Use all rows if max_rows_per_sheet is None, otherwise limit
                if max_rows_per_sheet is None:
                    data_df = df
                else:
                    data_df = df.head(max_rows_per_sheet)
                data_rows = data_df.fillna('').astype(str).apply(' | '.join, axis=1).tolist()
                text_parts.extend(data_rows)
                
                # Add summary if rows were truncated
                if max_rows_per_sheet and len(df) > max_rows_per_sheet:
                    text_parts.append(f"... ({len(df) - max_rows_per_sheet} more rows not shown)")
            else:
                text_parts.append("(Empty sheet)")
            
            text_parts.append("")  # Add blank line between sheets
            
        return "\n".join(text_parts), len(sheets)
    except Exception as e:
        logger.error(f"Error extracting text from Excel: {e}")
        raise

def extract_text_from_word(content_bytes):
    """Extract text from Word documents"""
    try:
        doc = Document(io.BytesIO(content_bytes))
        text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                text_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(text_parts), len(doc.paragraphs)
    except Exception as e:
        logger.error(f"Error extracting text from Word: {e}")
        raise

def extract_text_from_powerpoint(content_bytes):
    """Extract text from PowerPoint presentations"""
    try:
        prs = Presentation(io.BytesIO(content_bytes))
        text_parts = []
        
        for slide_num, slide in enumerate(prs.slides, 1):
            text_parts.append(f"=== Slide {slide_num} ===")
            
            # Extract text from shapes
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text.strip())
                
                # Extract text from tables
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip():
                            text_parts.append(row_text)
            
            text_parts.append("")  # Add blank line between slides
        
        return "\n".join(text_parts), len(prs.slides)
    except Exception as e:
        logger.error(f"Error extracting text from PowerPoint: {e}")
        raise

# Configuration
FILE_TYPES = {
    'text': ['.txt', '.csv', '.json', '.xml', '.html', '.md', '.js', '.css', '.py'],
    'pdf': ['.pdf'],
    'excel': ['.xlsx', '.xls'],
    'word': ['.docx', '.doc'],
    'powerpoint': ['.pptx', '.ppt']
}

def _list_folder_contents(folder_path: Optional[str] = None, item_type: str = "all") -> List[Dict[str, Any]]:
    """List contents of a folder using Graph API"""
    try:
        if not folder_path or folder_path == "":
            endpoint = f"sites/{SITE_ID}/drives/{DRIVE_ID}/root/children"
        else:
            endpoint = f"sites/{SITE_ID}/drives/{DRIVE_ID}/root:/{folder_path}:/children"
        
        response = make_graph_request("GET", endpoint)
        
        if response and response.status_code == 200:
            items = response.json().get('value', [])
            
            # Filter by item type
            if item_type == "folders":
                items = [item for item in items if 'folder' in item]
            elif item_type == "files":
                items = [item for item in items if 'file' in item]
            
            # Convert to standard format
            result = []
            for item in items:
                item_info = {
                    "name": item.get('name'),
                    "id": item.get('id'),
                    "url": item.get('webUrl'),
                    "created": item.get('createdDateTime'),
                    "modified": item.get('lastModifiedDateTime'),
                    "type": "folder" if 'folder' in item else "file"
                }
                
                if 'file' in item:
                    item_info["size"] = item.get('size', 0)
                    item_info["download_url"] = item.get('@microsoft.graph.downloadUrl')
                
                result.append(item_info)
            
            return result
        else:
            logger.error(f"Failed to list folder contents {folder_path}")
            return []
            
    except Exception as e:
        logger.error(f"Error listing folder contents {folder_path}: {e}")
        return []

def list_folders(parent_folder: Optional[str] = None) -> List[Dict[str, Any]]:
    """List folders using Graph API"""
    logger.info(f"Listing folders in {parent_folder or 'root directory'}")
    return _list_folder_contents(parent_folder, "folders")

def list_documents(folder_name: str) -> List[Dict[str, Any]]:
    """List documents using Graph API"""
    logger.info(f"Listing documents in folder: {folder_name}")
    return _list_folder_contents(folder_name, "files")

def _get_file_content_by_path(file_path: str) -> Optional[bytes]:
    """Download file content using Graph API"""
    try:
        endpoint = f"sites/{SITE_ID}/drives/{DRIVE_ID}/root:/{file_path}"
        response = make_graph_request("GET", endpoint)
        
        if not response or response.status_code != 200:
            return None
        
        file_info = response.json()
        download_url = file_info.get('@microsoft.graph.downloadUrl')
        
        if not download_url:
            return None
        
        # Download file content
        content_response = requests.get(download_url)
        return content_response.content if content_response.status_code == 200 else None
            
    except Exception as e:
        logger.error(f"Error downloading file {file_path}: {e}")
        return None

def get_document_content(folder_name: str, file_name: str) -> dict:
    """Get document content using Graph API with text extraction support"""
    file_path = f"{folder_name}/{file_name}".strip('/')
    logger.info(f"Getting content for file: {file_path}")
    
    try:
        content_bytes = _get_file_content_by_path(file_path)
        
        if not content_bytes:
            return {"name": file_name, "error": "Failed to download file content"}
        
        # Determine file type
        lower_name = file_name.lower()
        file_type = next((t for t, exts in FILE_TYPES.items() if any(lower_name.endswith(ext) for ext in exts)), 'binary')
        
        if file_type == 'pdf':
            try:
                text, pages = extract_text_from_pdf(content_bytes)
                return {"name": file_name, "content_type": "text", "content": text, "original_type": "pdf", "page_count": pages, "size": len(content_bytes)}
            except Exception as e:
                logger.warning(f"PDF processing failed: {e}")
                return {"name": file_name, "content_type": "binary", "content_base64": base64.b64encode(content_bytes).decode(), "original_type": "pdf", "size": len(content_bytes)}
        
        if file_type == 'excel':
            try:
                text, sheets = extract_text_from_excel(content_bytes)
                return {"name": file_name, "content_type": "text", "content": text, "original_type": "excel", "sheet_count": sheets, "size": len(content_bytes)}
            except Exception as e:
                logger.warning(f"Excel processing failed: {e}")
                return {"name": file_name, "content_type": "binary", "content_base64": base64.b64encode(content_bytes).decode(), "original_type": "excel", "size": len(content_bytes)}
        
        if file_type == 'word':
            try:
                text, paragraphs = extract_text_from_word(content_bytes)
                return {"name": file_name, "content_type": "text", "content": text, "original_type": "word", "paragraph_count": paragraphs, "size": len(content_bytes)}
            except Exception as e:
                logger.warning(f"Word processing failed: {e}")
                return {"name": file_name, "content_type": "binary", "content_base64": base64.b64encode(content_bytes).decode(), "original_type": "word", "size": len(content_bytes)}
        
        if file_type == 'powerpoint':
            try:
                text, slides = extract_text_from_powerpoint(content_bytes)
                return {"name": file_name, "content_type": "text", "content": text, "original_type": "powerpoint", "slide_count": slides, "size": len(content_bytes)}
            except Exception as e:
                logger.warning(f"PowerPoint processing failed: {e}")
                return {"name": file_name, "content_type": "binary", "content_base64": base64.b64encode(content_bytes).decode(), "original_type": "powerpoint", "size": len(content_bytes)}
        
        if file_type == 'text':
            try:
                return {"name": file_name, "content_type": "text", "content": content_bytes.decode('utf-8'), "size": len(content_bytes)}
            except UnicodeDecodeError:
                pass
        
        return {"name": file_name, "content_type": "binary", "content_base64": base64.b64encode(content_bytes).decode(), "size": len(content_bytes)}
        
    except Exception as e:
        logger.error(f"Failed to get document content for {file_path}: {e}")
        return {"name": file_name, "error": f"Failed to get document content: {str(e)}"}

def download_document(folder_name: str, file_name: str, local_path: str) -> Dict[str, Any]:
    """Download a document from SharePoint to local filesystem using Graph API"""
    file_path = f"{folder_name}/{file_name}".strip('/')
    logger.info(f"Downloading file: {file_path} to {local_path}")
    try:
        # Get file content using existing function
        content_bytes = _get_file_content_by_path(file_path)
        
        if not content_bytes:
            return {"success": False, "message": "Failed to download file content"}
        
        # Ensure local directory exists
        os.makedirs(os.path.dirname(local_path), exist_ok=True)
        
        # Write file to local filesystem
        with open(local_path, 'wb') as f:
            f.write(content_bytes)
        
        logger.info(f"Successfully downloaded {file_name} to {local_path}")
        return {
            "success": True,
            "message": f"File {file_name} downloaded successfully",
            "local_path": local_path,
            "size": len(content_bytes)
        }
        
    except Exception as e:
        logger.error(f"Failed to download document {file_path}: {e}")
        return {"success": False, "message": f"Failed to download document: {str(e)}"}