"""
PowerPoint helper for creating professional presentation reports
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
import pandas as pd
from datetime import datetime
from typing import Dict, List, Optional
from .common import logger
import io
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import numpy as np
from PIL import Image
import base64

class PowerPointHelper:
    """Helper class for creating professional PowerPoint presentations"""
    
    # Professional color scheme
    COLORS = {
        'primary': RGBColor(0, 112, 192),      # Professional blue
        'secondary': RGBColor(68, 114, 196),   # Light blue
        'accent': RGBColor(237, 125, 49),      # Orange
        'success': RGBColor(112, 173, 71),     # Green
        'header_bg': RGBColor(68, 114, 196),   # Header background
        'text': RGBColor(64, 64, 64),          # Dark gray text
        'light_gray': RGBColor(217, 217, 217), # Light gray
    }
    
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(7.5)
        
    def create_title_slide(self, title: str, subtitle: str = None):
        """Create professional title slide with centered title and date"""
        slide_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(slide_layout)
        
        # Add title - centered
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(2.5), Inches(9), Inches(2)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.word_wrap = True
        title_para = title_frame.paragraphs[0]
        title_para.alignment = PP_ALIGN.CENTER
        title_para.font.size = Pt(36)
        title_para.font.bold = True
        title_para.font.color.rgb = self.COLORS['primary']
        
        # Add subtitle/date - centered
        if subtitle is None:
            subtitle = datetime.now().strftime("%B %d, %Y")
        
        subtitle_box = slide.shapes.add_textbox(
            Inches(1), Inches(4.5), Inches(8), Inches(0.8)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle
        subtitle_para = subtitle_frame.paragraphs[0]
        subtitle_para.alignment = PP_ALIGN.CENTER
        subtitle_para.font.size = Pt(24)
        subtitle_para.font.color.rgb = self.COLORS['text']
        
        return slide
    
    def create_content_slide(self, title: str):
        """Create slide with professional header"""
        slide_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(slide_layout)
        
        # Add header background with color
        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(10), Inches(0.8)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = self.COLORS['header_bg']
        header_shape.line.fill.background()
        
        # Add title text in header
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.15), Inches(9), Inches(0.5)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(28)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(255, 255, 255)  # White text
        
        return slide
    
    def create_matplotlib_chart(self, chart_data: Dict, chart_type: str, width_inches: float = 5.0, height_inches: float = 4.0):
        """Create chart using matplotlib and return as image bytes"""
        try:
            labels = chart_data.get('labels', [])
            values = chart_data.get('values', [])
            
            if not labels or not values:
                return None
            
            # Set up matplotlib with professional styling
            plt.style.use('default')
            fig, ax = plt.subplots(figsize=(width_inches, height_inches), dpi=150)
            
            # Professional color palette
            colors = ['#0070C0', '#44729C', '#ED7D31', '#70AD47', '#FFC000', '#5B9BD5', '#A5A5A5', '#264478']
            
            if chart_type in ['pie']:
                # Create pie chart with legend
                wedges, texts, autotexts = ax.pie(values, labels=labels, autopct='%1.1f%%', 
                                                 colors=colors[:len(values)], startangle=90)
                ax.set_title('')
                # Position legend below chart
                ax.legend(wedges, labels, loc='upper center', bbox_to_anchor=(0.5, -0.1), ncol=2)
                
            else:
                # Create bar chart without legend
                bars = ax.bar(labels, values, color=colors[:len(values)])
                ax.set_ylabel('Count')
                ax.set_title('')
                
                # Handle x-axis labels intelligently
                max_label_length = max(len(str(label)) for label in labels) if labels else 0
                
                if len(labels) > 4 or max_label_length > 12:
                    # Rotate labels for better readability
                    plt.xticks(rotation=45, ha='right', fontsize=9)
                    # Add more bottom margin for rotated labels
                    plt.subplots_adjust(bottom=0.2)
                elif max_label_length > 8:
                    # Smaller font for medium-length labels
                    plt.xticks(fontsize=9)
                else:
                    # Normal font for short labels
                    plt.xticks(fontsize=10)
                
                # Add value labels on top of bars
                for bar in bars:
                    height = bar.get_height()
                    ax.text(bar.get_x() + bar.get_width()/2., height + max(values)*0.01,
                           f'{height:.1f}', ha='center', va='bottom', fontsize=9)
            
            # Clean up the chart appearance
            ax.spines['top'].set_visible(False)
            ax.spines['right'].set_visible(False)
            ax.grid(True, alpha=0.3, axis='y')
            
            # Tight layout to minimize whitespace
            plt.tight_layout()
            
            # Save to bytes
            img_buffer = io.BytesIO()
            plt.savefig(img_buffer, format='png', bbox_inches='tight', dpi=150, 
                       facecolor='white', edgecolor='none')
            img_buffer.seek(0)
            plt.close(fig)  # Clean up memory
            
            return img_buffer.getvalue()
            
        except Exception as e:
            logger.error(f"Error creating matplotlib chart: {e}")
            return None

    def add_chart_to_slide(self, slide, chart_data: Dict, chart_type: str, 
                          left: float, top: float, width: float, height: float):
        """Add chart to slide using matplotlib for better control"""
        try:
            # Create chart image using matplotlib
            chart_image_bytes = self.create_matplotlib_chart(chart_data, chart_type, width, height)
            
            if chart_image_bytes:
                # Add image to slide
                img_stream = io.BytesIO(chart_image_bytes)
                picture = slide.shapes.add_picture(
                    img_stream, Inches(left), Inches(top), 
                    Inches(width), Inches(height)
                )
                return picture
            else:
                return None
                
        except Exception as e:
            logger.error(f"Error adding chart to slide: {e}")
            return None
    
    def add_bullet_points(self, slide, insights: List[str], 
                         left: float, top: float, width: float, height: float):
        """Add bullet points to slide"""
        text_box = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        
        for i, insight in enumerate(insights):
            if i == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()
            
            p.text = insight
            p.level = 0
            p.font.size = Pt(14)
            p.font.color.rgb = self.COLORS['text']
            p.space_before = Pt(6)
            p.space_after = Pt(6)
        
        return text_box
    
    def add_table(self, slide, data: List[List[str]], 
                 left: float, top: float, width: float, height: float,
                 header_row: bool = True):
        """Add formatted table to slide"""
        rows = len(data)
        cols = len(data[0]) if data else 0
        
        if rows == 0 or cols == 0:
            return None
        
        table = slide.shapes.add_table(
            rows, cols, Inches(left), Inches(top),
            Inches(width), Inches(height)
        ).table
        
        # Populate table
        for i, row_data in enumerate(data):
            for j, cell_value in enumerate(row_data):
                cell = table.cell(i, j)
                cell.text = str(cell_value)
                cell.text_frame.paragraphs[0].font.size = Pt(11)
                
                # Format header row
                if i == 0 and header_row:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.COLORS['header_bg']
                    cell.text_frame.paragraphs[0].font.bold = True
                    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        return table
    
    def create_chart_slide(self, title: str, chart_data: Dict, chart_type: str, 
                          insights: List[str]):
        """Create slide with chart and insights"""
        slide = self.create_content_slide(title)
        
        # Add chart on left side with more space for legend
        self.add_chart_to_slide(
            slide, chart_data, chart_type,
            left=0.5, top=1.2, width=5.0, height=4.8
        )
        
        # Add insights on right side (avoid overlap)
        if insights:
            # Add "Key Insights" label
            label_box = slide.shapes.add_textbox(
                Inches(6.5), Inches(1.2), Inches(3), Inches(0.4)
            )
            label_frame = label_box.text_frame
            label_frame.text = "Key Insights"
            label_para = label_frame.paragraphs[0]
            label_para.font.size = Pt(18)
            label_para.font.bold = True
            label_para.font.color.rgb = self.COLORS['primary']
            
            # Add bullet points
            self.add_bullet_points(
                slide, insights,
                left=6.5, top=1.8, width=3, height=4
            )
        
        return slide
    
    def create_data_definitions_slide(self, columns: Dict[str, str]):
        """Create slide explaining data column definitions"""
        slide = self.create_content_slide("Data Column Definitions")
        
        # Convert to table format
        table_data = [["Column Name", "Definition"]]
        for col_name, definition in columns.items():
            table_data.append([col_name, definition])
        
        # Add table
        self.add_table(
            slide, table_data,
            left=0.5, top=1.2, width=9, height=5.5,
            header_row=True
        )
        
        return slide
    
    def create_insight_slide(self, title: str, insight_data: Dict):
        """
        Create flexible insight slide based on insight type
        
        Supported insight types:
        - 'metric': Key metric with value and context
        - 'comparison': Side-by-side comparison
        - 'recommendation': Actionable recommendation with rationale
        - 'analysis': Detailed analysis with bullet points
        - 'chart_with_analysis': Chart with detailed insights
        - 'table': Data table with insights
        - 'two_column': Two-column layout with sections
        """
        insight_type = insight_data.get('type', 'analysis')
        
        if insight_type == 'metric':
            return self._create_metric_slide(title, insight_data)
        elif insight_type == 'comparison':
            return self._create_comparison_slide(title, insight_data)
        elif insight_type == 'recommendation':
            return self._create_recommendation_slide(title, insight_data)
        elif insight_type == 'chart_with_analysis':
            return self._create_chart_analysis_slide(title, insight_data)
        elif insight_type == 'table':
            return self._create_table_slide(title, insight_data)
        elif insight_type == 'two_column':
            return self._create_two_column_slide(title, insight_data)
        else:  # Default to 'analysis'
            return self._create_analysis_slide(title, insight_data)
    
    def _create_metric_slide(self, title: str, data: Dict):
        """Create slide highlighting a key metric"""
        slide = self.create_content_slide(title)
        
        # Large metric value in center
        metric_value = data.get('value', '')
        metric_label = data.get('label', '')
        
        # Metric value - large and centered
        value_box = slide.shapes.add_textbox(
            Inches(2), Inches(2), Inches(6), Inches(1.5)
        )
        value_frame = value_box.text_frame
        value_frame.text = str(metric_value)
        value_para = value_frame.paragraphs[0]
        value_para.alignment = PP_ALIGN.CENTER
        value_para.font.size = Pt(72)
        value_para.font.bold = True
        value_para.font.color.rgb = self.COLORS['primary']
        
        # Metric label
        label_box = slide.shapes.add_textbox(
            Inches(2), Inches(3.5), Inches(6), Inches(0.5)
        )
        label_frame = label_box.text_frame
        label_frame.text = metric_label
        label_para = label_frame.paragraphs[0]
        label_para.alignment = PP_ALIGN.CENTER
        label_para.font.size = Pt(24)
        label_para.font.color.rgb = self.COLORS['text']
        
        # Context/insights below
        context = data.get('context', [])
        if context:
            self.add_bullet_points(slide, context, left=2, top=4.5, width=6, height=2)
        
        return slide
    
    def _create_comparison_slide(self, title: str, data: Dict):
        """Create slide with side-by-side comparison"""
        slide = self.create_content_slide(title)
        
        left_content = data.get('left', {})
        right_content = data.get('right', {})
        
        # Left side
        left_title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.2), Inches(4), Inches(0.5)
        )
        left_title_frame = left_title_box.text_frame
        left_title_frame.text = left_content.get('title', '')
        left_title_para = left_title_frame.paragraphs[0]
        left_title_para.font.size = Pt(20)
        left_title_para.font.bold = True
        left_title_para.font.color.rgb = self.COLORS['primary']
        
        if left_content.get('points'):
            self.add_bullet_points(slide, left_content['points'], 
                                 left=0.5, top=1.8, width=4, height=4.5)
        
        # Right side
        right_title_box = slide.shapes.add_textbox(
            Inches(5.5), Inches(1.2), Inches(4), Inches(0.5)
        )
        right_title_frame = right_title_box.text_frame
        right_title_frame.text = right_content.get('title', '')
        right_title_para = right_title_frame.paragraphs[0]
        right_title_para.font.size = Pt(20)
        right_title_para.font.bold = True
        right_title_para.font.color.rgb = self.COLORS['accent']
        
        if right_content.get('points'):
            self.add_bullet_points(slide, right_content['points'],
                                 left=5.5, top=1.8, width=4, height=4.5)
        
        return slide
    
    def _create_recommendation_slide(self, title: str, data: Dict):
        """Create slide with actionable recommendation"""
        slide = self.create_content_slide(title)
        
        # Recommendation box with colored background
        rec_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1), Inches(1.5), Inches(8), Inches(1.2)
        )
        rec_shape.fill.solid()
        rec_shape.fill.fore_color.rgb = self.COLORS['success']
        rec_shape.line.fill.background()
        
        # Recommendation text
        rec_text = data.get('recommendation', '')
        rec_box = slide.shapes.add_textbox(
            Inches(1.2), Inches(1.7), Inches(7.6), Inches(0.8)
        )
        rec_frame = rec_box.text_frame
        rec_frame.text = rec_text
        rec_frame.word_wrap = True
        rec_para = rec_frame.paragraphs[0]
        rec_para.font.size = Pt(18)
        rec_para.font.bold = True
        rec_para.font.color.rgb = RGBColor(255, 255, 255)
        
        # Rationale section
        rationale = data.get('rationale', [])
        if rationale:
            rationale_label = slide.shapes.add_textbox(
                Inches(1), Inches(3), Inches(8), Inches(0.4)
            )
            rationale_label.text_frame.text = "Rationale:"
            rationale_label.text_frame.paragraphs[0].font.size = Pt(16)
            rationale_label.text_frame.paragraphs[0].font.bold = True
            rationale_label.text_frame.paragraphs[0].font.color.rgb = self.COLORS['primary']
            
            self.add_bullet_points(slide, rationale, left=1, top=3.5, width=8, height=2.5)
        
        return slide
    
    def _create_analysis_slide(self, title: str, data: Dict):
        """Create slide with detailed analysis points"""
        slide = self.create_content_slide(title)
        
        # Main content
        content = data.get('content', [])
        if content:
            self.add_bullet_points(slide, content, left=0.8, top=1.2, width=8.4, height=5.5)
        
        # Optional footer note
        footer = data.get('footer', '')
        if footer:
            footer_box = slide.shapes.add_textbox(
                Inches(0.8), Inches(6.8), Inches(8.4), Inches(0.4)
            )
            footer_frame = footer_box.text_frame
            footer_frame.text = footer
            footer_para = footer_frame.paragraphs[0]
            footer_para.font.size = Pt(10)
            footer_para.font.italic = True
            footer_para.font.color.rgb = self.COLORS['text']
        
        return slide
    
    def _create_chart_analysis_slide(self, title: str, data: Dict):
        """Create slide with chart and detailed analysis"""
        slide = self.create_content_slide(title)
        
        # Chart on left
        chart_data = data.get('chart_data', {})
        chart_type = data.get('chart_type', 'bar')
        if chart_data:
            self.add_chart_to_slide(
                slide, chart_data, chart_type,
                left=0.5, top=1.2, width=5.0, height=5.0
            )
        
        # Analysis on right
        analysis = data.get('analysis', [])
        if analysis:
            self.add_bullet_points(slide, analysis, left=6, top=1.2, width=3.5, height=5.5)
        
        return slide
    
    def _create_table_slide(self, title: str, data: Dict):
        """Create slide with data table and insights"""
        slide = self.create_content_slide(title)
        
        # Table
        table_data = data.get('table_data', [])
        if table_data:
            self.add_table(
                slide, table_data,
                left=0.5, top=1.2, width=9, height=4,
                header_row=data.get('header_row', True)
            )
        
        # Insights below table
        insights = data.get('insights', [])
        if insights:
            self.add_bullet_points(slide, insights, left=0.5, top=5.5, width=9, height=1.5)
        
        return slide
    
    def _create_two_column_slide(self, title: str, data: Dict):
        """Create slide with two-column layout"""
        slide = self.create_content_slide(title)
        
        sections = data.get('sections', [])
        
        # Left column
        if len(sections) > 0:
            left_section = sections[0]
            left_title = slide.shapes.add_textbox(
                Inches(0.5), Inches(1.2), Inches(4.5), Inches(0.4)
            )
            left_title.text_frame.text = left_section.get('title', '')
            left_title.text_frame.paragraphs[0].font.size = Pt(18)
            left_title.text_frame.paragraphs[0].font.bold = True
            left_title.text_frame.paragraphs[0].font.color.rgb = self.COLORS['primary']
            
            if left_section.get('content'):
                self.add_bullet_points(slide, left_section['content'],
                                     left=0.5, top=1.7, width=4.5, height=5)
        
        # Right column
        if len(sections) > 1:
            right_section = sections[1]
            right_title = slide.shapes.add_textbox(
                Inches(5.5), Inches(1.2), Inches(4), Inches(0.4)
            )
            right_title.text_frame.text = right_section.get('title', '')
            right_title.text_frame.paragraphs[0].font.size = Pt(18)
            right_title.text_frame.paragraphs[0].font.bold = True
            right_title.text_frame.paragraphs[0].font.color.rgb = self.COLORS['primary']
            
            if right_section.get('content'):
                self.add_bullet_points(slide, right_section['content'],
                                     left=5.5, top=1.7, width=4, height=5)
        
        return slide
    
    def save_to_bytes(self) -> bytes:
        """Save presentation to bytes"""
        pptx_buffer = io.BytesIO()
        self.prs.save(pptx_buffer)
        pptx_buffer.seek(0)
        return pptx_buffer.getvalue()

# Global instance
powerpoint_helper = PowerPointHelper()
